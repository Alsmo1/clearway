# server/excel_export.py
"""
Builds one organized, human-readable workbook from the current database
state. This is what gets saved to your local backups folder on manual export
and on the
monthly automatic backup.
"""
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from . import db

HEADER_FILL = PatternFill(start_color="FF5B7F66", end_color="FF5B7F66", fill_type="solid")
HEADER_FONT = Font(color="FFFFFFFF", bold=True)


def _style_header(ws, row_idx: int, ncols: int):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row_idx, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[row_idx].height = 20


def _autowidth(ws):
    for col_cells in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
        col_letter = get_column_letter(col_cells[0].column)
        ws.column_dimensions[col_letter].width = min(length + 2, 42)


def build_workbook() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)

    habits = db.all_active("habits")
    logs = db.all_active("habit_logs")
    todos = db.all_active("todos")
    expenses = db.all_active("expenses")
    goals = db.all_active("goals")

    # --- Summary sheet ---------------------------------------------------
    summary = wb.create_sheet("Summary")
    summary.append(["ClearWay — Export Summary"])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append([f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"])
    summary.append([])
    total_spent = sum(e["amount"] for e in expenses)
    open_todos = sum(1 for t in todos if not t["completed"])
    active_habits = sum(1 for h in habits if not h["archived"])
    summary.append(["Metric", "Value"])
    _style_header(summary, summary.max_row, 2)
    summary.append(["Active habits", active_habits])
    summary.append(["Habit check-ins logged", sum(1 for l in logs if l["completed"])])
    summary.append(["Open to-dos", open_todos])
    summary.append(["Completed to-dos", sum(1 for t in todos if t["completed"])])
    summary.append(["Total expenses logged", len(expenses)])
    summary.append(["Total amount spent", round(total_spent, 2)])
    summary.append(["Active goals", len(goals)])
    _autowidth(summary)

    # --- Habits sheet ------------------------------------------------------
    hs = wb.create_sheet("Habits")
    hs.append(["Name", "Frequency", "Target / week", "Archived", "Created"])
    _style_header(hs, 1, 5)
    for h in habits:
        hs.append([h["name"], h["frequency"], h["target_per_week"], "Yes" if h["archived"] else "No", (h["created_at"] or "")[:10]])
    _autowidth(hs)

    # --- Habit check-ins sheet -------------------------------------------
    ls = wb.create_sheet("Habit Check-ins")
    ls.append(["Habit", "Date", "Completed"])
    _style_header(ls, 1, 3)
    name_by_id = {h["id"]: h["name"] for h in habits}
    for l in sorted(logs, key=lambda x: x["date"], reverse=True):
        ls.append([name_by_id.get(l["habit_id"], "(deleted habit)"), l["date"], "Yes" if l["completed"] else "No"])
    _autowidth(ls)

    # --- To-Dos sheet -------------------------------------------------------
    ts = wb.create_sheet("To-Dos")
    ts.append(["Title", "Notes", "Due date", "Priority", "Completed", "Completed at"])
    _style_header(ts, 1, 6)
    for t in todos:
        ts.append([
            t["title"], t.get("notes") or "", t.get("due_date") or "", t["priority"],
            "Yes" if t["completed"] else "No", (t.get("completed_at") or "")[:10],
        ])
    _autowidth(ts)

    # --- Expenses sheet ------------------------------------------------
    es = wb.create_sheet("Expenses")
    es.append(["Date", "Category", "Amount", "Note"])
    _style_header(es, 1, 4)
    for e in sorted(expenses, key=lambda x: x["date"], reverse=True):
        es.append([e["date"], e["category"], e["amount"], e.get("note") or ""])
    for row in es.iter_rows(min_row=2, min_col=3, max_col=3):
        for cell in row:
            cell.number_format = "#,##0.00"

    by_category: dict[str, float] = {}
    for e in expenses:
        by_category[e["category"]] = by_category.get(e["category"], 0) + e["amount"]
    es.append([])
    cat_row = es.max_row + 1
    es.append(["Category breakdown"])
    es[f"A{cat_row}"].font = Font(bold=True)
    for cat, amt in by_category.items():
        es.append(["", cat, round(amt, 2)])
    _autowidth(es)

    # --- Goals sheet -------------------------------------------------------
    gs = wb.create_sheet("Goals")
    gs.append(["Title", "Progress", "Target", "Unit", "Deadline"])
    _style_header(gs, 1, 5)
    for g in goals:
        gs.append([g["title"], g["current_value"], g["target_value"], g["unit"], g.get("deadline") or ""])
    _autowidth(gs)

    return wb
